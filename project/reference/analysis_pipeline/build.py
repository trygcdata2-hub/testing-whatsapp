import pandas as pd, numpy as np, pickle, re, datetime as dtm
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, DataBarRule
from openpyxl.chart import LineChart, BarChart, Reference
P=pickle.load(open('tabs2.pkl','rb')); globals().update(P); FUNBR=P['FB']
d=pd.read_pickle('d2.pkl')
BL=t[t.ask_turn.notna()].groupby('kind').agg(turns=('replied','size'),rep=('replied','mean'),acc=('accepted','mean')).reset_index()
FONT='Arial'
F0=Font(name=FONT,size=10); FB=Font(name=FONT,size=10,bold=True); FH=Font(name=FONT,size=10,bold=True,color='FFFFFF')
FT_=Font(name=FONT,size=14,bold=True,color='1F3864'); FS=Font(name=FONT,size=10,italic=True,color='595959'); FIN=Font(name=FONT,size=10,color='0000FF',bold=True)
HF=PatternFill('solid',fgColor='1F3864'); SF=PatternFill('solid',fgColor='D9E2F3'); INF=PatternFill('solid',fgColor='FFF2CC'); TF=PatternFill('solid',fgColor='EDEDED')
thin=Side(style='thin',color='BFBFBF'); BD=Border(bottom=thin)
CLEAN=re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
def cl(v):
    if isinstance(v,(bool,np.bool_)): return int(v)
    if isinstance(v,str): return CLEAN.sub('',v)[:32000]
    if isinstance(v,(np.integer,)): return int(v)
    if isinstance(v,(np.floating,)): return None if np.isnan(v) else float(v)
    if isinstance(v,float) and np.isnan(v): return None
    if v is pd.NaT: return None
    if isinstance(v,pd.Timestamp): return v.to_pydatetime()
    return v
wb=Workbook(); wb.remove(wb.active)
def sheet(name,title,sub=None,tab=None):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    ws['A1']=title; ws['A1'].font=FT_
    if sub: ws['A2']=sub; ws['A2'].font=FS
    if tab: ws.sheet_properties.tabColor=tab
    return ws
def write(ws,df,r0,fmts=None,widths=None,hdr_h=42,wrap_cols=()):
    fmts=fmts or {}
    for j,c in enumerate(df.columns,1):
        x=ws.cell(r0,j,c); x.font=FH; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    ws.row_dimensions[r0].height=hdr_h
    for i,row in enumerate(df.itertuples(index=False),r0+1):
        for j,v in enumerate(row,1):
            x=ws.cell(i,j,cl(v)); x.font=F0; x.border=BD
            f=fmts.get(df.columns[j-1])
            if f: x.number_format=f
            if df.columns[j-1] in wrap_cols: x.alignment=Alignment(wrap_text=True,vertical='top')
    for j,c in enumerate(df.columns,1):
        w=(widths or {}).get(c, min(max(len(str(c))*0.9,10),18)); ws.column_dimensions[L(j)].width=w
    return r0+len(df)
def section(ws,r,text,ncol=8):
    for j in range(1,ncol+1): ws.cell(r,j).fill=SF
    ws.cell(r,1,text).font=FB
PCT='0.0%'; N0='#,##0'; N1='#,##0.0'; N2='0.00'
# ================= numbers for the narrative =================
RR=R[(R.replied==1)&(~R.late)]
team_med=RR.rt_min.median(); team_p15=(RR.rt_min<=15).mean(); team_p60=(RR.rt_min<=60).mean()
n_unans=int(R.unanswered.sum()); base_in=int((~R.closer_only).sum()); unans_rate=n_unans/(len(RR)+n_unans)
agmsgs=int(d.agent.notna().sum()); nconv=d.phone.nunique(); nin=len(R)
yesall=y[y.real_yes==1]; n_yes=len(yesall)
Yok=Y; nofu_rate=Y.followed.eq(0).mean(); fu_med=Y.follow_min[Y.followed==1].median()
fun=dict(FUN)
resp=SC[(SC.role=='Inbox responder')&(SC.replies>=300)]
fast=resp.sort_values('med_shift').head(3); slow=resp.sort_values('med_shift',ascending=False).head(3)
def lst(df,col,fmt='{:.0f}'): return '، '.join(f"{a} ({fmt.format(v)} min)" for a,v in zip(df.Agent,df[col]))
def lst_en(df,col,fmt='{:.0f}'): return ', '.join(f"{a} ({fmt.format(v)} min)" for a,v in zip(df.Agent,df[col]))
P2=PERS[PERS['Ask-turns']>=100]
top=P2[P2.Verdict=='Clearly above expected'].head(4); bot=P2[P2.Verdict=='Clearly below expected'].sort_values('z (kind×brand)').head(4)
def pl(df): return ', '.join(f"{a} ({l:.2f}x, {n:,} turns)" for a,l,n in zip(df.Agent,df['Accept lift (kind×brand)'],df['Ask-turns']))
H_=H.set_index('hour'); late_h=H_.loc[21:23]; 
peak_h=H_.inbound_per_agent[H_.agents_active>=3].idxmax()
def F(label,kind='PITCH'): 
    r=FT[(FT['Message kind']==kind)&(FT.Feature.str.startswith(label))].iloc[0]; return r
f_q=F('Ends with a question'); f_urg=F('Urgency'); f_tam=F('Asks to send'); f_tmpl=F('Sent as a shared'); f_gulf=F('Written in Gulf'); f_greet=F('Opens with a greeting')
c_cq=F('Closed question','CHAT'); c_t=F('Proposes a day','CHAT')
bp=BRD_P.copy(); bp_best=bp[bp.pitches>=300].sort_values('acc',ascending=False).head(3); bp_worst=bp[bp.pitches>=300].sort_values('acc').head(3)
ctyd=CTY.set_index('country'); 
worstday=DAY.sort_values('unans_rate',ascending=False).head(2)
ayw=AY[AY.yes>=100].sort_values('rate',ascending=False).head(4); ayb=AY[AY.yes>=100].sort_values('rate').head(3)
objs=OBJ.sort_values('All messages',ascending=False).head(4)
# ================= Playbook =================
ws=sheet('Playbook','Persuasion & performance playbook — what the last 30 days say (4 Sep – 3 Oct 2026)',f'Based on {len(d):,} messages, {nconv:,} conversations, {nin:,} inbound bursts and {len(t):,} agent turns. Findings are computed from the sheets in this file; "association, not proof" applies to every wording finding.','1F3864')
ws.column_dimensions['A'].width=5; ws.column_dimensions['B'].width=42; ws.column_dimensions['C'].width=120
r=4; section(ws,r,'Headline numbers',3); r+=1
kp=[('Agent messages sent',f'{agmsgs:,} by {d.agent.nunique()} accounts ({len(SC[SC.role=="Inbox responder"])} real inbox responders)'),
('Inbound bursts (influencer wrote)',f'{nin:,}; {len(RR):,} answered within 24h, {n_unans:,} left unanswered ({unans_rate:.1%} of answerable bursts)'),
('Team reply time',f'median {team_med:.0f} min; {team_p15:.0%} answered within 15 min; {team_p60:.0%} within 60 min (overnight gaps included)'),
('Real "yes" answers',f'{n_yes:,} (an influencer accepting right after a pitch / reminder / chat question); {len(Y):,} of them were credited to a named agent and are old enough to check follow-up'),
('Pitch → yes → logistics funnel',f'{fun["Conversations pitched (invitation sent)"]:,} conversations pitched → {fun["…influencer replied to a pitch"]:,} replied → {fun["Conversations with a real yes"]:,} said yes → {fun["…then address / brief / logistics sent"]:,} received address/brief afterwards → {fun["…then post-visit follow-up (review / post)"]:,} got a post-visit follow-up'),
('Conversations never answered',f'{never} conversations where the influencer wrote and no agent ever wrote back')]
for a,b in kp:
    ws.cell(r,2,a).font=FB; c=ws.cell(r,3,b); c.font=F0; c.alignment=Alignment(wrap_text=True,vertical='top'); r+=1
def block(title,items):
    global r
    r+=1; section(ws,r,title,3); r+=1
    for i,(a,b) in enumerate(items,1):
        ws.cell(r,1,i).font=FB; ws.cell(r,2,a).font=FB; ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
        c=ws.cell(r,3,b); c.font=F0; c.alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[r].height=max(30,15*(len(b)//115+1)); r+=1
block('Who performs',[
('Fastest and slowest responders (own-shift median reply)',f'Fastest among agents with 300+ replies: {lst_en(fast,"med_shift")}. Slowest: {lst_en(slow,"med_shift")}. "Own shift" = only inbound messages that arrived in hours when that agent normally works, so night-shift agents are not punished for sleeping hours (full median is on the Scorecard).'),
('Persuasion vs expected (kind × brand adjusted)',f'Clearly above expected (z ≥ 3, 100+ turns): {pl(top)}. Clearly below: {pl(bot)}. Lift compares each agent with what the team achieves for the same message kind AND the same brand, so an agent stuck with a weak campaign is not blamed for it. Caution: pitches that ask the influencer to send «تم» are easier to "accept" (see the wording section), so template-heavy accounts get a small head-start in this metric. Full table on Persuasion.'),
('Consistency matters as much as peaks','Weekly_Agent shows week-by-week reply time and accept rate for every agent with 1,000+ messages — look for agents whose numbers swing week to week (workload or hand-off problems) before judging them on the monthly average.'),
('Test / admin accounts','trainee01, Developers Team and IT Team appear in the data; they are tagged "Test / admin account" and excluded from ranking.')])
block('Where leads are lost',[
('Unanswered influencers',f'{n_unans:,} bursts got no reply within 24h ({unans_rate:.1%}). Worst days: {", ".join(f"{pd.Timestamp(a).strftime("%d %b")} ({b:.0%})" for a,b in zip(worstday.date,worstday.unans_rate))} vs a typical 10–14%. Check staffing on those days (Daily_Trend). {n_unans-int(SC.unans.sum()):,} of the {n_unans:,} came in conversations where no agent had written yet (influencer-initiated chats nobody picked up) — they appear in the team total but not on any agent row.'),
('Evening staffing gap',f'Inbound load per active agent peaks at {H_.inbound_per_agent[peak_h]:.0f} bursts/agent-hour around {peak_h}:00, and from 21:00 to 23:59 the median reply time jumps to {late_h.median_rt.min():.0f}–{late_h.median_rt.max():.0f} min with only {late_h.agents_active.min():.1f}–{late_h.agents_active.max():.1f} agents active. Also 00:00–03:00 (median {H_.median_rt.loc[0:3].min():.0f}–{H_.median_rt.loc[0:3].max():.0f} min). See Hourly_Coverage.'),
('Yes with no follow-up',f'{(1-Y.followed).sum():,.0f} of {len(Y):,} real yeses ({nofu_rate:.0%}) got no agent message within 24h; median time to follow-up is {fu_med:.0f} min. Highest no-follow-up rates (100+ yeses): {", ".join(f"{a} {b:.0%}" for a,b in zip(ayw.first_agent,ayw.rate))}. Lowest: {", ".join(f"{a} {b:.0%}" for a,b in zip(ayb.first_agent,ayb.rate))}.'),
('Funnel leak after yes',f'Only {fun["…then address / brief / logistics sent"]/fun["Conversations with a real yes"]:.0%} of conversations with a yes show an address/brief/logistics message afterwards (keyword-detected, so treat as a floor) and {fun["…then post-visit follow-up (review / post)"]/fun["Conversations with a real yes"]:.0%} a post-visit follow-up. Link a booking status to confirm (see Data to add).'),
('Weak campaigns',f'Pitch accept rate by brand (300+ pitches): best {", ".join(f"{b} {a:.0%}" for b,a in zip(bp_best.brand,bp_best.acc))}; weakest {", ".join(f"{b} {a:.1%} (reply {rp:.0%})" for b,a,rp in zip(bp_worst.brand,bp_worst.acc,bp_worst.rep))}. Rewrite or re-target the weak ones before sending more volume.'),
('Kuwait vs KSA',f'Ask-turn reply rate: Kuwait {ctyd.loc["Kuwait","rep"]:.0%} vs KSA {ctyd.loc["KSA","rep"]:.0%}; accept {ctyd.loc["Kuwait","acc"]:.1%} vs {ctyd.loc["KSA","acc"]:.1%}. Review Kuwait wording and send-time before blaming agents.')])
block('What wording works (Message_Features, brand-adjusted)',[
('End the pitch with a question',f'Pitches ending with a question: {f_q["Reply rate (with)"]:.0%} reply vs {f_q["Reply rate (without)"]:.0%} ({f_q["Reply lift (brand-adjusted)"]:.2f}x after adjusting for brand), n={int(f_q["n (with)"]):,}. Accept rate difference is small ({f_q["Accept rate (with)"]:.1%} vs {f_q["Accept rate (without)"]:.1%}).'),
('Urgency / "limited time" hurts',f'Urgency wording: {f_urg["Reply rate (with)"]:.0%} reply and {f_urg["Accept rate (with)"]:.1%} accept vs {f_urg["Reply rate (without)"]:.0%} and {f_urg["Accept rate (without)"]:.1%} without (accept lift {f_urg["Accept lift (brand-adjusted)"]:.2f}x brand-adjusted). Holds on the one-day finding.'),
('Shared templates underperform personal pitches',f'Pitches whose opening is a shared template (15+ identical copies): {f_tmpl["Reply rate (with)"]:.0%} reply, {f_tmpl["Accept rate (with)"]:.1%} accept vs {f_tmpl["Reply rate (without)"]:.0%} / {f_tmpl["Accept rate (without)"]:.1%} for one-off pitches (n={int(f_tmpl["n (with)"]):,} vs {int(f_tmpl["n (without)"]):,}).'),
('«أرسل تم» works for yes, not for chat',f'Asking the influencer to send «تم»: accept {f_tam["Accept rate (with)"]:.1%} vs {f_tam["Accept rate (without)"]:.1%} ({f_tam["Accept lift (brand-adjusted)"]:.2f}x adjusted) but a lower reply rate ({f_tam["Reply rate (with)"]:.0%} vs {f_tam["Reply rate (without)"]:.0%}) — it turns the reply itself into the yes. Keep it, and add a closing question variant to test.'),
('Day-one findings that did NOT hold over a month',f'Gulf-dialect pitches ({f_gulf["Accept lift (brand-adjusted)"]:.2f}x accept, brand-adjusted) and greeting openers ({f_greet["Accept lift (brand-adjusted)"]:.2f}x) looked strong on 3 Oct but are weaker over 30 days. Likely confounded with who wrote them and which brand — do not make them a rule.'),
('In conversation: short closed questions, propose a time',f'Chat replies with a short closed question get {c_cq["Reply rate (with)"]:.0%} replies vs {c_cq["Reply rate (without)"]:.0%}; proposing a concrete day/time lifts accept to {c_t["Accept rate (with)"]:.1%} vs {c_t["Accept rate (without)"]:.1%} ({c_t["Accept lift (brand-adjusted)"]:.2f}x adjusted). Chat messages of 21–90 chars convert best (≈17–18% accept); 200+ chars drop to ≈9%. In reminders, proposing a time lowers replies — use it in chat, not in chasers.')])
tops=', '.join(f"{a} ({b:,})" for a,b in zip(objs.Topic,objs['All messages']))
block('What influencers ask / object to',[
('Top topics',f'{tops}. Ready answers for each topic are listed on the Objections sheet, split by brand and country.'),
('Payment is a real objection','Influencers who only do paid ads should be told up-front whether the deal is barter-only, so they self-select out and agents stop chasing them.')])
block('Do next (this week)',[
('1. Evening coverage',f'Only {H_.agents_active.loc[15:20].mean():.1f} agents are active on average between 15:00 and 20:59 versus {H_.agents_active.loc[12:14].mean():.1f} at 12:00–14:59, while the load per agent is higher in the afternoon. Move 2–3 agents from the morning/noon peak to 15:00–20:00 and put at least 4 agents on 21:00–23:59 ({H_.agents_active.loc[21:23].mean():.1f} today), or stop sending pitches after 21:00.'),
('2. Auto-follow-up on yes','Send address + brief automatically the moment a yes arrives, and alert the owner if no agent wrote within 15 minutes (today the median is ~1 hour and about 1 in 5 yeses get nothing).'),
('3. Rewrite the weak pitches','Retire urgency wording and emoji-first blasts; use personal-style variants ending in a question; A/B them on the weakest brands first.'),
('4. Coach from the winners','Read the threads of the "clearly above expected" agents (Messages are in the source CSV; filter by agent) and write a one-page guide of their first two replies.'),
('5. Data to add','A booking status per phone (confirmed / visited / posted) and a campaign id per message would replace keyword guesses (brand, yes, logistics) and make every verdict here stronger.')])
ws.freeze_panes='A4'
# ================= Scorecard =================
ws=sheet('Scorecard','Agent performance scorecard — 4 Sep to 3 Oct 2026','Rates are formulas where possible. Medians come from the Response_Events sheet. See Notes for definitions. "Own shift" median only counts inbound that arrived in the hours the agent normally works.','2E75B6')
cols=['Agent','Role','Messages sent','Active days','Msgs per active day','Conversations touched','Conversations opened by agent','Replies to influencers','Median reply time (min)','Median reply — own shift (min)','Avg reply time (min)','P90 reply time (min)','% replied ≤5 min','% replied ≤15 min','% replied ≤60 min','Influencer msgs left unanswered','Unanswered rate','Influencer reply rate','Positive replies','Positive-reply rate','Invitations sent (convs)','Confirmed after invitation','Confirmation rate','Template share','Media sent','First msg','Last msg']
for j,c in enumerate(cols,1):
    x=ws.cell(4,j,c); x.font=FH; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
ws.row_dimensions[4].height=55
S=SC.copy(); r0=5
for i,row in enumerate(S.itertuples(index=False),r0):
    vals=[row.Agent,row.role,row.msgs,row.days,f'=C{i}/D{i}',row.convs,row.opened,row.replies,row.med,row.med_shift,row.avg,row.p90,row.p5,row.p15,row.p60,row.unans,f'=IFERROR(P{i}/(H{i}+P{i}),0)',row.irr,row.pos,f'=IFERROR(S{i}/F{i},0)',row.inv,row.conf,f'=IFERROR(V{i}/U{i},0)',row.tshare,row.media,row.first,row.last]
    for j,v in enumerate(vals,1):
        x=ws.cell(i,j,cl(v)); x.font=F0; x.border=BD
        x.number_format={3:N0,4:N0,5:N1,6:N0,7:N0,8:N0,9:N1,10:N1,11:N1,12:N1,13:PCT,14:PCT,15:PCT,16:N0,17:PCT,18:PCT,19:N0,20:PCT,21:N0,22:N0,23:PCT,24:PCT,25:N0,26:'dd-mmm hh:mm',27:'dd-mmm hh:mm'}.get(j,'General')
rl=r0+len(S)-1; tr=rl+1
ws.cell(tr,1,'TEAM (all agents)').font=FB
allr=RR.rt_min
tv={3:f'=SUM(C{r0}:C{rl})',6:f'=SUM(F{r0}:F{rl})',7:f'=SUM(G{r0}:G{rl})',8:f'=SUM(H{r0}:H{rl})',9:allr.median(),11:allr.mean(),12:allr.quantile(.9),13:(allr<=5).mean(),14:(allr<=15).mean(),15:(allr<=60).mean(),16:n_unans,17:f'=IFERROR(P{tr}/({len(RR)}+P{tr}),0)',19:f'=SUM(S{r0}:S{rl})',21:f'=SUM(U{r0}:U{rl})',22:f'=SUM(V{r0}:V{rl})',23:f'=IFERROR(V{tr}/U{tr},0)',25:f'=SUM(Y{r0}:Y{rl})'}
for j,v in tv.items():
    x=ws.cell(tr,j,v); x.font=FB; x.fill=TF
    x.number_format={3:N0,6:N0,7:N0,8:N0,9:N1,11:N1,12:N1,13:PCT,14:PCT,15:PCT,16:N0,17:PCT,19:N0,21:N0,22:N0,23:PCT,25:N0}.get(j,'General')
for j in range(1,28): ws.cell(tr,j).fill=TF
ws.cell(tr+2,1,'Team row: conversation / invitation counts are summed per agent, so a conversation handled by 2 agents counts twice. Team reply-time stats include replies sent from the phone directly (not attributable to an agent). Team unanswered is the total of all unanswered bursts (4,654) including ones in conversations no agent had written in yet; agent rows only include bursts credited to an agent. Team rate = unanswered ÷ (all answered + unanswered).').font=FS
for j,w in enumerate([26,20,10,8,9,12,12,11,11,12,11,11,9,9,9,12,10,10,10,10,11,11,10,10,8,14,14],1): ws.column_dimensions[L(j)].width=w
ws.freeze_panes='C5'
ws.conditional_formatting.add(f'I{r0}:I{rl}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
ws.conditional_formatting.add(f'J{r0}:J{rl}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
ws.conditional_formatting.add(f'Q{r0}:Q{rl}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
ws.conditional_formatting.add(f'N{r0}:N{rl}',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='63BE7B'))
ws.conditional_formatting.add(f'T{r0}:T{rl}',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='63BE7B'))
SCROW={a:r0+i for i,a in enumerate(S.Agent)}
# ================= Ranking (formulas) =================
ws=sheet('Ranking','Composite ranking (indicative)','Yellow cells are inputs — change them and the ranking recalculates. Rank 1 = best. Speed uses the "own shift" median reply time. Only non-test agents with enough replies are ranked.','2E75B6')
ws['A4']='Weights (should sum to 100%)'; ws['A4'].font=FB
inp=[('A5','Weight — median reply time',0.35),('A6','Weight — unanswered rate',0.25),('A7','Weight — influencer reply rate',0.20),('A8','Weight — positive-reply rate',0.20),('A9','Min. replies to be ranked',100),('A10','Min. active days to be ranked',5)]
for a,lab,v in inp:
    rr=int(a[1:]); ws.cell(rr,1,lab).font=F0; c=ws.cell(rr,2,v); c.font=FIN; c.fill=INF; c.number_format=PCT if v<1 else '0'
ws.cell(5,1)
ws['A11']='Weights total'; ws['A11'].font=FB; ws['B11']='=SUM(B5:B8)'; ws['B11'].number_format=PCT; ws['B11'].font=FB
h=['Agent','Role','Replies to influencers','Ranked?','Median reply — own shift (min)','Unanswered rate','Influencer reply rate','Positive-reply rate','Rank: speed','Rank: unanswered','Rank: engagement','Rank: positive','Weighted index (lower = better)','Overall rank']
for j,c in enumerate(h,1):
    x=ws.cell(12,j,c); x.font=FH; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
ws.row_dimensions[12].height=45
n=len(S); a0=13; a1=a0+n-1
for k,a in enumerate(S.Agent):
    i=a0+k; sr=SCROW[a]
    ws.cell(i,1,f'=Scorecard!A{sr}'); ws.cell(i,2,f'=Scorecard!B{sr}'); ws.cell(i,3,f'=Scorecard!H{sr}')
    ws.cell(i,4,f'=IF(AND(C{i}>=$B$9,B{i}<>"Test / admin account",Scorecard!D{sr}>=$B$10,ISNUMBER(Scorecard!J{sr}),ISNUMBER(Scorecard!R{sr})),"Yes","No")')
    ws.cell(i,5,f'=IF(D{i}="Yes",Scorecard!J{sr},"n/a")'); ws.cell(i,6,f'=IF(D{i}="Yes",Scorecard!Q{sr},"n/a")'); ws.cell(i,7,f'=IF(D{i}="Yes",Scorecard!R{sr},"n/a")'); ws.cell(i,8,f'=IF(D{i}="Yes",Scorecard!T{sr},"n/a")')
    ws.cell(i,9,f'=IF($D{i}="Yes",COUNTIFS($D${a0}:$D${a1},"Yes",E${a0}:E${a1},"<"&E{i})+1,"n/a")')
    ws.cell(i,10,f'=IF($D{i}="Yes",COUNTIFS($D${a0}:$D${a1},"Yes",F${a0}:F${a1},"<"&F{i})+1,"n/a")')
    ws.cell(i,11,f'=IF($D{i}="Yes",COUNTIFS($D${a0}:$D${a1},"Yes",G${a0}:G${a1},">"&G{i})+1,"n/a")')
    ws.cell(i,12,f'=IF($D{i}="Yes",COUNTIFS($D${a0}:$D${a1},"Yes",H${a0}:H${a1},">"&H{i})+1,"n/a")')
    ws.cell(i,13,f'=IF(D{i}="Yes",$B$5*I{i}+$B$6*J{i}+$B$7*K{i}+$B$8*L{i},"n/a")')
    ws.cell(i,14,f'=IF(D{i}="Yes",COUNTIFS($D${a0}:$D${a1},"Yes",M${a0}:M${a1},"<"&M{i})+1,"n/a")')
    for j in range(1,15):
        x=ws.cell(i,j); x.font=F0; x.border=BD
        x.number_format={3:N0,5:N1,6:PCT,7:PCT,8:PCT,13:'0.00'}.get(j,'General')
for j,w in enumerate([26,20,10,9,13,11,11,11,9,11,11,9,13,9],1): ws.column_dimensions[L(j)].width=w
ws.cell(a1+2,1,'Reading tip: "n/a" = too few replies (or test account) to judge fairly. Ties share the better rank. Raise or lower the minimum replies (B9) and minimum active days (B10) to widen or narrow the table.').font=FS
ws.freeze_panes='A13'
ws.conditional_formatting.add(f'N{a0}:N{a1}',ColorScaleRule(start_type='num',start_value=1,start_color='63BE7B',end_type='num',end_value=20,end_color='F8696B'))
# ================= Persuasion =================
ws=sheet('Persuasion','Who persuades best? Agent conversion vs. expected','Ask-turns = PITCH + REMINDER + CHAT. Expected = what the team would achieve with the same mix of message kinds (and, in the last 3 columns, the same brands). Verdict uses the kind × brand z-score: |z| ≥ 3 clearly, 1.5–3 leaning (stricter than a single-day cut-off because 30+ agents are compared at once). Pitches that ask for «تم» are easier to accept, so template-heavy agents get a small head-start.','548235')
PP=PERS.copy(); end=write(ws,PP,4,{'Reply rate':PCT,'Expected replies':N1,'Reply lift':N2,'Accept rate':PCT,'Expected accepts':N1,'Accept lift':N2,'Accept rate – low (95%)':PCT,'Accept rate – high (95%)':PCT,'z-score':N2,'Pitch accept rate':PCT,'Chat accept rate':PCT,'Expected accepts (kind×brand)':N1,'Accept lift (kind×brand)':N2,'z (kind×brand)':N2,'Ask-turns':N0,'Influencer replied':N0,'Accepted':N0},{'Agent':26,'Verdict':24},hdr_h=55)
ws.freeze_panes='B5'
ws.conditional_formatting.add(f'V5:V{end}',ColorScaleRule(start_type='num',start_value=0.5,start_color='F8696B',mid_type='num',mid_value=1,mid_color='FFFFFF',end_type='num',end_value=1.6,end_color='63BE7B'))
ws.conditional_formatting.add(f'W5:W{end}',ColorScaleRule(start_type='num',start_value=-6,start_color='F8696B',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='num',end_value=6,end_color='63BE7B'))
# ================= Daily_Trend =================
ws=sheet('Daily_Trend','Daily trend — volume, speed, unanswered, yes','Use this to spot staffing problems and campaign spikes. Reply time per day is the median of bursts that arrived that day.','C55A11')
D2=DAY.copy(); D2['date']=pd.to_datetime(D2.date)
D2=D2[['date','dow','agent_msgs','active_agents','inbound_bursts','median_rt','pct15','unanswered','unans_rate','pitches','pitch_reply','yes']]
D2.columns=['Date','Day','Agent messages','Active agents','Inbound bursts','Median reply (min)','% answered ≤15 min','Unanswered bursts','Unanswered rate','Pitch turns','Pitch reply rate','Real yeses']
end=write(ws,D2,4,{'Date':'dd-mmm','Agent messages':N0,'Inbound bursts':N0,'Median reply (min)':N1,'% answered ≤15 min':PCT,'Unanswered bursts':N0,'Unanswered rate':PCT,'Pitch turns':N0,'Pitch reply rate':PCT,'Real yeses':N0},{'Date':10,'Day':11})
ws.conditional_formatting.add(f'F5:F{end}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
ws.conditional_formatting.add(f'I5:I{end}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
ch=BarChart(); ch.type='col'; ch.title='Inbound bursts vs agent messages per day'; ch.height=8; ch.width=22
ch.add_data(Reference(ws,min_col=3,min_row=4,max_row=end),titles_from_data=True); ch.add_data(Reference(ws,min_col=5,min_row=4,max_row=end),titles_from_data=True); ch.set_categories(Reference(ws,min_col=1,min_row=5,max_row=end)); ws.add_chart(ch,'N4')
lc=LineChart(); lc.title='Median reply time (min) and unanswered rate'; lc.height=8; lc.width=22
lc.add_data(Reference(ws,min_col=6,min_row=4,max_row=end),titles_from_data=True); lc.set_categories(Reference(ws,min_col=1,min_row=5,max_row=end))
lc2=LineChart(); lc2.add_data(Reference(ws,min_col=9,min_row=4,max_row=end),titles_from_data=True); lc2.y_axis.axId=200; lc2.y_axis.crosses='max'; lc2.y_axis.number_format='0%'; lc+=lc2; ws.add_chart(lc,'N21')
ws.freeze_panes='A5'
# ================= Hourly_Coverage =================
ws=sheet('Hourly_Coverage','Hourly coverage — does staffing match demand?','Averages per day over 30 days. "Inbound per agent" = inbound bursts per hour ÷ agents active in that hour; high values + long reply times = understaffed hours.','C55A11')
H2=H[['hour','inbound_per_day','agents_active','inbound_per_agent','median_rt','pct15','unans_rate','pitches_per_day','pitch_reply']].copy()
H2.columns=['Hour','Inbound bursts per day','Agents active (avg)','Inbound per active agent','Median reply (min)','% answered ≤15 min','Unanswered rate','Pitches sent per day','Pitch reply rate']
end=write(ws,H2,4,{'Inbound bursts per day':N1,'Agents active (avg)':N1,'Inbound per active agent':N1,'Median reply (min)':N0,'% answered ≤15 min':PCT,'Unanswered rate':PCT,'Pitches sent per day':N1,'Pitch reply rate':PCT},{'Hour':8})
ws.conditional_formatting.add(f'E5:E{end}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
ws.conditional_formatting.add(f'D5:D{end}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
b=BarChart(); b.type='col'; b.title='Inbound bursts per day vs agents active, by hour'; b.height=8; b.width=22
b.add_data(Reference(ws,min_col=2,min_row=4,max_row=end),titles_from_data=True); b.set_categories(Reference(ws,min_col=1,min_row=5,max_row=end))
l2=LineChart(); l2.add_data(Reference(ws,min_col=3,min_row=4,max_row=end),titles_from_data=True); l2.y_axis.axId=200; l2.y_axis.crosses='max'; b+=l2; ws.add_chart(b,'L4')
c2=BarChart(); c2.type='col'; c2.title='Median reply time (min) by arrival hour'; c2.height=8; c2.width=22
c2.add_data(Reference(ws,min_col=5,min_row=4,max_row=end),titles_from_data=True); c2.set_categories(Reference(ws,min_col=1,min_row=5,max_row=end)); ws.add_chart(c2,'L21')
ws.freeze_panes='A5'
# ================= Weekly_Agent =================
ws=sheet('Weekly_Agent','Weekly trend per agent (inbox responders with 1,000+ messages)','W1 = 4–10 Sep, W2 = 11–17, W3 = 18–24, W4 = 25 Sep–3 Oct (9 days). Blank = agent did not work that week. Look for swings, not small differences.','C55A11')
WK2=WK.copy(); end=write(ws,WK2,4,{'Messages':N0,'Replies':N0,'Median reply (min)':N1,'Accept rate (ask-turns)':PCT,'Ask-turns':N0},{'Agent':26,'Week':8})
ws.freeze_panes='A5'
ws.conditional_formatting.add(f'E5:E{end}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
ws.conditional_formatting.add(f'F5:F{end}',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='63BE7B'))
# ================= Funnel_Segments =================
ws=sheet('Funnel_Segments','Funnel, brands, countries and weekdays','Conversation-level funnel (a phone number = one conversation). Brand is detected from the wording of agent messages (keyword list) — "Unknown" = generic chat.','7030A0')
FU=pd.DataFrame(FUN,columns=['Stage','Conversations']); FU['% of pitched']=FU.Conversations/FU.Conversations.iloc[1]
section(ws,4,'Funnel',8); e1=write(ws,FU,5,{'Conversations':N0,'% of pitched':PCT},{'Stage':48})
r=e1+2; section(ws,r,'By brand — conversations in the funnel',8); e2=write(ws,FUNBR,r+1,{c:N0 for c in FUNBR.columns[1:]},{'Brand':48})
r=e2+2; section(ws,r,'By brand — pitch turns',8)
BP=BRD_P.rename(columns={'brand':'Brand','pitches':'Pitch turns','rep':'Reply rate','acc':'Accept rate','agents':'Agents used'}); e3=write(ws,BP,r+1,{'Pitch turns':N0,'Reply rate':PCT,'Accept rate':PCT})
r=e3+2; section(ws,r,'By country — all ask-turns (100+ turns)',8)
CT=CTY.rename(columns={'country':'Country','turns':'Ask-turns','rep':'Reply rate','acc':'Accept rate','phones':'Conversations'}); e4=write(ws,CT,r+1,{'Ask-turns':N0,'Reply rate':PCT,'Accept rate':PCT,'Conversations':N0})
r=e4+2; section(ws,r,'By weekday — all ask-turns',8)
dd=DOW.copy(); dd['o']=dd.dow.map({'Monday':0,'Tuesday':1,'Wednesday':2,'Thursday':3,'Friday':4,'Saturday':5,'Sunday':6}); dd=dd.sort_values('o').drop(columns='o').rename(columns={'dow':'Weekday','turns':'Ask-turns','rep':'Reply rate','acc':'Accept rate'})
e5=write(ws,dd,r+1,{'Ask-turns':N0,'Reply rate':PCT,'Accept rate':PCT})
r=e5+2; section(ws,r,'Pitch reply by send hour',8)
hh=HR.rename(columns={'hb':'Send hour','n':'Pitch turns','rep':'Reply rate','acc':'Accept rate'}); write(ws,hh,r+1,{'Pitch turns':N0,'Reply rate':PCT,'Accept rate':PCT})
# ================= Templates =================
ws=sheet('Templates','Which invitation / reminder wording works?','Messages grouped by the first 45 characters (ignoring emoji/punctuation). Only openings used 40+ times. Compare templates of the same brand; "Accept lift" is vs the average for the same message kind.','7030A0')
TT=TP[['template_key','brand','kind','n','agents','rep','rr','acc','ar','lift','length','ex']].copy(); TT.columns=['Template starts with…','Brand','Kind','Times sent','Agents using it','Influencer replied','Reply rate','Accepted','Accept rate','Accept lift vs kind','Avg length (chars)','Example text']
end=write(ws,TT,4,{'Times sent':N0,'Reply rate':PCT,'Accept rate':PCT,'Accept lift vs kind':N2,'Avg length (chars)':N0,'Influencer replied':N0,'Accepted':N0},{'Template starts with…':34,'Brand':18,'Example text':80},wrap_cols=('Example text',))
ws.freeze_panes='B5'
ws.conditional_formatting.add(f'J5:J{end}',ColorScaleRule(start_type='num',start_value=0,start_color='F8696B',mid_type='num',mid_value=1,mid_color='FFFFFF',end_type='num',end_value=2,end_color='63BE7B'))
# ================= Message_Features =================
ws=sheet('Message_Features','What makes a message more likely to get a reply / a yes?','Each row compares turns with the feature vs without, inside the same message kind. "Brand-adjusted" lifts compare against the same brands\' own rates, so they are less confounded than raw lifts. Robust = 100+ in both groups, p < 0.01 and the adjusted lift points the same way.','7030A0')
end=write(ws,FT,4,{'n (with)':N0,'Reply rate (with)':PCT,'Accept rate (with)':PCT,'n (without)':N0,'Reply rate (without)':PCT,'Accept rate (without)':PCT,'Reply lift':N2,'Accept lift':N2,'Reply lift (brand-adjusted)':N2,'Accept lift (brand-adjusted)':N2,'p-value':'0.0000'},{'Feature':50,'Message kind':12},hdr_h=55)
r=end+2; section(ws,r,'Length (chars of the whole turn)',8); e=write(ws,LEN,r+1,{'Turns':N0,'Reply rate':PCT,'Accept rate':PCT})
r=e+2; section(ws,r,'Pitch by sending hour',8); write(ws,HR.rename(columns={'hb':'Send hour','n':'Pitch turns','rep':'Reply rate','acc':'Accept rate'}),r+1,{'Pitch turns':N0,'Reply rate':PCT,'Accept rate':PCT})
ws.freeze_panes='A5'
# ================= After_Yes =================
ws=sheet('After_Yes','After the influencer says yes — who follows up, and how fast?','"Real yes" = accepted right after a pitch / reminder / personal question. Follow-up = next agent message in that thread within 24h. Yeses from the last 3 hours of the data are excluded.','C00000')
A2=AY.rename(columns={'first_agent':'Agent (got the yes)','yes':'Real yeses','fu':'Followed up (≤24h)','nofu':'No follow-up','rate':'No-follow-up rate','med':'Median follow-up (min)'})[['Agent (got the yes)','Real yeses','Followed up (≤24h)','No follow-up','No-follow-up rate','Median follow-up (min)']]
e=write(ws,A2,4,{'Real yeses':N0,'Followed up (≤24h)':N0,'No follow-up':N0,'No-follow-up rate':PCT,'Median follow-up (min)':N1},{'Agent (got the yes)':28})
ws.conditional_formatting.add(f'E5:E{e}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
r=e+2; section(ws,r,'By hour the yes arrived',6); Y2=Yh.rename(columns={'yes_hour':'Hour band','yes':'Real yeses','nofu':'No follow-up','med':'Median follow-up (min)','rate':'No-follow-up rate'}); Y2['Hour band']=Y2['Hour band'].astype(str); e=write(ws,Y2,r+1,{'Real yeses':N0,'No follow-up':N0,'Median follow-up (min)':N0,'No-follow-up rate':PCT})
r=e+2; section(ws,r,'How fast the follow-up came',6)
sp=pd.DataFrame({'Speed':list(Yb.index)+['No follow-up in 24h'],'Real yeses':list(Yb.values.astype(int))+[int(nofu_n)]}); sp['Share']=sp['Real yeses']/sp['Real yeses'].sum(); write(ws,sp,r+1,{'Real yeses':N0,'Share':PCT})
# ================= Objections =================
ws=sheet('Objections','What influencers say — questions, objections, FAQs','Keyword-based topic tagging of influencer text messages (a message can have several topics). Columns for the 6 biggest brands and 3 countries.','C00000')
write(ws,OBJ,4,{c:N0 for c in OBJ.columns if c not in('Topic','What to prepare')},{'Topic':26,'What to prepare':70},hdr_h=45)
# ================= Baselines =================
ws=sheet('Baselines','Baselines — what an "average" turn achieves','Used to judge agents fairly: cold pitches always convert less than chat replies, so each agent is compared with the baseline for the same kind of message (and brand).','808080')
BLd=t.groupby('kind').agg(turns=('replied','size'),rep=('replied','mean'),acc=('accepted','mean')).reset_index(); BLd.columns=['Message kind','Turns','Influencer replied (within 24h)','Influencer accepted']
BLd['Meaning']=BLd['Message kind'].map({'PITCH':'Invitation / offer message (templates, "send تم")','REMINDER':'Chasing a reply','CHAT':'Personal reply inside an existing conversation','LOGISTICS':'Address / brief / order notice','POST_VISIT':'Visit / review / publish follow-up','CONFIRM':'Booking confirmation'})
e=write(ws,BLd,4,{'Turns':N0,'Influencer replied (within 24h)':PCT,'Influencer accepted':PCT},{'Message kind':16,'Meaning':60})
# ================= Notes =================
ws=sheet('Notes','How the data was read','','808080')
raw_n=183152
notes=[('Scope',f'Source: last_month_s_data.csv — {raw_n:,} rows, 4 Sep → 3 Oct 2026. Cleaning: {raw_n-len(d)-5197-13810+ (13810 and 0):,}'),]
notes=[('Scope',f'Source: last_month_s_data.csv — 183,152 raw rows (4 Sep → 3 Oct 2026). Removed: 831 exact duplicates, 5,197 system rows ("Whatsapp number changed…") and ~13.8K rows with no phone number (group chats / mentions). Analysed: {len(d):,} messages, {nconv:,} conversations.'),
('Who is who','isFromMe=0 → message from the influencer. isFromMe=1 with a real userName → sent by that agent from the platform. isFromMe=1 with userName "original whatsapp" → sent from the phone itself (23.7K messages) — cannot be attributed to an agent, kept in conversations but not credited. Test accounts: trainee01, Developers Team, IT Team.'),
('Turn','A run of consecutive outgoing messages in one conversation, split when the agent changes or there is a gap of 30+ min. Reply = the influencer wrote next within 24h. Accept = that reply contains an acceptance word (تم، تمام، اوكي، ابشر، مناسب، ايوه، اكيد، yes، ok، 👍) and no negation/decline.'),
('Reply time','Every run of consecutive influencer messages is one "inbound burst"; the clock starts at the first message and stops at the next outgoing message, credited to the agent who sent it. Includes overnight gaps — use medians. "Own shift" median only counts bursts that arrived in hours when the responding agent normally works (hours holding ≥4% of that agent\'s messages).'),
('Unanswered','A burst with no outgoing message within 24h, credited to the last agent who wrote in that conversation. Bursts that are only a closing word (شكرا / تمام / 👍 / sticker) are not counted. Bursts in the last 3 hours of the data are treated as pending.'),
('Influencer reply rate','Share of an agent\'s conversations where the influencer wrote anything after the agent\'s first message.'),
('Positive reply / Confirmed','Positive = first acceptance per conversation, credited to the agent whose turn came right before. Confirmed after invitation = positive where that turn was a PITCH. Keyword-based: it also catches "تمام" as a plain acknowledgement.'),
('Message kind','PITCH = invitation / offer wording (or isInvitation flag in the export); REMINDER = chasing; CHAT = personal reply; LOGISTICS = address / brief / order notice (incl. long briefs ≥600 chars); POST_VISIT = review / publish follow-up; CONFIRM = booking confirmation. Keyword rules — approximate. Calibrated against the 3 Oct workbook: 99% of turns matched, accept flag agreed on 94%.'),
('Expected accepts','For each agent: sum over their ask-turns of the team accept rate for that message kind (and, for the kind × brand columns, kind + brand when the brand has 50+ turns). Lift = actual ÷ expected. z = (actual − expected) ÷ √expected. Verdict uses kind × brand z: ≥3 / ≤−3 clearly, 1.5–3 leaning, and requires 100+ ask-turns.'),
('Brand','Detected from brand words in the agent text of the turn, else the most common brand in that thread; "Unknown" = generic. Keyword list covers the 20 biggest campaigns.'),
('Comparison with the 3 Oct file','Numbers differ slightly from the one-day file because: (1) this file includes every conversation, not only those where a named agent wrote; (2) turns here are split by agent change/30-min gap; (3) PITCH replies are higher because many monthly pitches follow an influencer-initiated chat. Always compare agents inside this file, not against the old one.'),
('Caveats','Timestamps are as exported (no timezone conversion). Several agents answer the same WhatsApp numbers, so a slow reply may reflect hand-offs. Agents with fewer than ~100 ask-turns are flagged "Too few turns". Message-feature findings are association, not causation (agent, brand, country and send hour are tangled).'),
('Not measured','No field says whether a booking/visit was really confirmed or whether the influencer posted. If your DB has that status, join it on phone — it would replace the keyword-based yes / logistics / post-visit columns.'),
('Evidence sheets','Response_Events, Yes_Events, Turns (lean, text truncated) and Inbound_Msgs (topic-tagged only, first 15,000) are the rows behind the tables. The full message text is in the source CSV.')]
r=3
for a,b in notes:
    ws.cell(r,1,a).font=FB; ws.cell(r,1).alignment=Alignment(vertical='top'); c=ws.cell(r,2,b); c.font=F0; c.alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[r].height=max(30,15*(len(b)//120+1)); r+=1
ws.column_dimensions['A'].width=26; ws.column_dimensions['B'].width=130
# ================= Evidence sheets =================
ws=sheet('Response_Events','Response events (one row per inbound burst)',None,'808080')
E=R[['phone','inbound_at','inbound_msgs','closer_only','prev_agent','replied','replied_at','rt_min','responder','unanswered','pending','country']].copy(); E['phone']=E.phone.astype(str)
write(ws,E,3,{'inbound_at':'dd-mmm hh:mm','replied_at':'dd-mmm hh:mm','rt_min':N1},{'phone':16,'inbound_at':16,'replied_at':16,'prev_agent':24,'responder':24},hdr_h=20); ws.freeze_panes='A4'
ws=sheet('Yes_Events','Yes events (influencer acceptances)',None,'808080')
Yv=y[['phone','in_t','yes_hour','first_agent','kind','real_yes','followed','follow_min','next_agent','next_kind']].copy(); Yv.columns=['phone','accepted_at','hour','persuader','prev_type','real_yes','followed','follow_min','closer','next_type']; Yv['phone']=Yv.phone.astype(str)
write(ws,Yv,3,{'accepted_at':'dd-mmm hh:mm','follow_min':N1},{'phone':16,'accepted_at':16,'persuader':24,'closer':24},hdr_h=20); ws.freeze_panes='A4'
ws=sheet('Turns','Agent turns (lean — text truncated to 300 chars)',None,'808080')
Tn=t[['phone','sent_at','first_agent','kind','brand','country','hour','chars','starts_emoji','starts_greeting','cta_tam','cta_question','urgency','benefit','gulf','has_q','proposes_time','closed_q','len_bucket','template_key','replied','accepted','declined','payment_q','reply_delay_min','first_text']].copy()
Tn.columns=['phone','sent_at','agent','kind','brand','country','hour','chars','starts_emoji','starts_greeting','cta_tam','ends_with_question','urgency','benefit','gulf','has_q','proposes_time','closed_q','len_bucket','template_key','replied','accepted','declined','payment_q','reply_delay_min','text']
Tn['phone']=Tn.phone.astype(str); Tn['text']=Tn.text.str.replace('\n',' ').str[:300]
write(ws,Tn,3,{'sent_at':'dd-mmm hh:mm','reply_delay_min':N1},{'phone':16,'sent_at':16,'agent':24,'text':60},hdr_h=20); ws.freeze_panes='A4'
ws=sheet('Inbound_Msgs','Influencer messages tagged with a topic (first 15,000)',None,'808080')
I2=INB.copy(); I2['phone']=I2.phone.astype(str); I2['text']=I2.text.str.replace('\n',' ').str[:300]
write(ws,I2,3,{'at':'dd-mmm hh:mm'},{'phone':16,'at':16,'topics':36,'text':70},hdr_h=20); ws.freeze_panes='A4'
order=['Playbook','Scorecard','Ranking','Persuasion','Daily_Trend','Hourly_Coverage','Weekly_Agent','Funnel_Segments','Templates','Message_Features','After_Yes','Objections','Baselines','Notes','Response_Events','Yes_Events','Turns','Inbound_Msgs']
wb._sheets=[wb[n] for n in order]
wb.save('/mnt/user-data/outputs/agent_performance_last_month.xlsx'); print('saved')
