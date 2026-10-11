"""
pCloud Campaign Finder - READ ONLY

اللوجيك:
نمشي على كل فولدر — في كل فولدر نشوف محتواه (الأبناء):
  CASE A: في الأبناء فولدر اسمه Coverage → الفولدر الحالي = حملة
  CASE B: في الأبناء فولدر اسمه @username، وجوا الـ @username في platform →
           الفولدر الحالي = حملة
  CASE C: مفيش علامة → نكمل recursively على الأبناء

مهم: لما نحدد الحملة نوقف — مش بنكمل جواها
"""

import os, re, pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIG
# ============================================================

BASE_ACTIVE = (
    r"P:\pCloud Drive\Operation - Head Office\Coverages-Final"
    r"\2026- Active Campaigns all offices\pCloud Active Campaigns"
)

REGIONS = {
    "1": {"name": "KSA", "path": r"4- KSA", "has_groups": True},
    "2": {"name": "UAE", "path": r"3- UAE",  "has_groups": False},
    "3": {"name": "KW",  "path": r"2- KW",   "has_groups": False},
    "4": {"name": "EGY", "path": r"1- EGY",  "has_groups": False},
}

MAX_DEPTH = 10

# ============================================================
# PATTERNS
# ============================================================

COVERAGE_RE = re.compile(r'^cov(er(age)?s?|r?a?g?e?s?|\.)?$', re.IGNORECASE)

# platform لوحده = الكلمة كلها بالظبط
PLATFORM_RE = re.compile(
    r'^(instagram|insta|ig|tiktok|tt|facebook|fb|snapchat|snap|twitter|x|youtube|yt)$',
    re.IGNORECASE
)

def is_coverage(name): return bool(COVERAGE_RE.match(name.strip()))
def is_platform(name): return bool(PLATFORM_RE.match(name.strip()))

# ============================================================
# HELPERS
# ============================================================

def get_subfolders(path):
    try:
        return sorted([n for n in os.listdir(path)
                       if os.path.isdir(os.path.join(path, n))])
    except:
        return []


def folder_contains_coverage(path):
    """هل الفولدر ده فيه Coverage مباشرة؟"""
    return any(is_coverage(n) for n in get_subfolders(path))


def folder_contains_platform_via_users(path):
    """
    هل الفولدر ده فيه يوزرات، وجوا أي يوزر فيه platform؟
    (الحالة: حملة → @user → instagram)
    """
    children = get_subfolders(path)
    if not children:
        return False
    # نشوف أول 5 أبناء بس عشان منبطئش
    for child in children[:5]:
        child_path = os.path.join(path, child)
        grandchildren = get_subfolders(child_path)
        if any(is_platform(g) for g in grandchildren):
            return True
    return False


# ============================================================
# RECURSIVE SCANNER
# ============================================================

def scan(current_path, company_path, depth, found, results):
    """
    في كل فولدر نشوف:
      1. هل فيه Coverage بين أبناءه؟ → الفولدر ده = حملة
      2. هل فيه يوزرات وجوا اليوزرات platforms؟ → الفولدر ده = حملة
      3. غير كده → نكمل recursively
    """
    if depth > MAX_DEPTH:
        return

    # ── CASE A: في الأبناء Coverage ──────────────────────────
    if folder_contains_coverage(current_path):
        campaign_path = current_path
        if campaign_path not in found:
            found.add(campaign_path)
            campaign_name = os.path.basename(campaign_path)

            # Brand = الفولدر اللي فوق الحملة — لو مختلف عن الشركة
            parent = os.path.dirname(campaign_path)
            brand  = os.path.basename(parent) if parent != company_path else ""

            results.append({
                "Campaign Name": campaign_name,
                "Brand":         brand,
                "Campaign Path": campaign_path,
            })
            print(f"  ✅ [Cov]  {campaign_name}  |  brand: {brand or '-'}")
        return  # وقفنا — مش بندخل جوا الحملة

    # ── CASE B: يوزرات جواها platforms ───────────────────────
    if folder_contains_platform_via_users(current_path):
        campaign_path = current_path
        if campaign_path not in found:
            found.add(campaign_path)
            campaign_name = os.path.basename(campaign_path)

            parent = os.path.dirname(campaign_path)
            brand  = os.path.basename(parent) if parent != company_path else ""

            results.append({
                "Campaign Name": campaign_name,
                "Brand":         brand,
                "Campaign Path": campaign_path,
            })
            print(f"  ✅ [Plat] {campaign_name}  |  brand: {brand or '-'}")
        return  # وقفنا

    # ── CASE C: نكمل recursively على الأبناء ─────────────────
    for child in get_subfolders(current_path):
        child_path = os.path.join(current_path, child)
        scan(child_path, company_path, depth + 1, found, results)


# ============================================================
# RUNNER
# ============================================================

def run(region_path, has_groups):
    print(f"\n📂 Scanning: {region_path}\n")
    all_results = []
    found       = set()

    if has_groups:
        for group in get_subfolders(region_path):
            group_path = os.path.join(region_path, group)
            print(f"  [{group}]")
            for company in get_subfolders(group_path):
                company_path = os.path.join(group_path, company)
                print(f"    {company}")
                scan(company_path, company_path, depth=1, found=found, results=all_results)
    else:
        for company in get_subfolders(region_path):
            company_path = os.path.join(region_path, company)
            print(f"  {company}")
            scan(company_path, company_path, depth=1, found=found, results=all_results)

    print(f"\n📊 Total campaigns: {len(all_results)}")
    return all_results


# ============================================================
# EXCEL
# ============================================================

def export(results, output_file, region_name):
    if not results:
        print("⚠️  No campaigns found.")
        return

    df = pd.DataFrame(results, columns=["Campaign Name", "Brand", "Campaign Path"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{region_name} Campaigns"
    ws.sheet_properties.tabColor = "4267B2"

    hdr_fill = PatternFill("solid", fgColor="4267B2")
    alt_fill = PatternFill("solid", fgColor="EEF2FF")
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    dat_font = Font(name="Arial", size=10)
    thin     = Side(style="thin", color="DEE2E6")
    border   = Border(bottom=thin, right=thin)

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = dat_font; cell.border = border
            cell.alignment = Alignment(vertical="center")
            if row_idx % 2 == 0: cell.fill = alt_fill
        ws.row_dimensions[row_idx].height = 16

    for col_idx in range(1, len(df.columns) + 1):
        data = df.iloc[:, col_idx-1].astype(str)
        w = max(data.str.len().max() if len(data) else 0, len(df.columns[col_idx-1])) + 4
        ws.column_dimensions[get_column_letter(col_idx)].width = min(w, 80)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(output_file)
    print(f"\n💾 Saved: {output_file}")


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print("  pCloud Campaign Finder  |  READ ONLY")
    print("=" * 60)

    print("\nاختار الـ Region:")
    for k, v in REGIONS.items(): print(f"  {k}. {v['name']}")
    print("  5. Custom path")

    choice = input("\nاختيارك: ").strip()
    matched = next((k for k, v in REGIONS.items()
                    if choice == k or choice.startswith(k+".") or
                    choice.upper() == v["name"].upper()), None)
    if matched: choice = matched

    if choice in REGIONS:
        r = REGIONS[choice]
        region_path = os.path.join(BASE_ACTIVE, r["path"])
        has_groups  = r["has_groups"]
        region_name = r["name"]
    elif choice.startswith("5"):
        region_path = input("Full path: ").strip().strip('"')
        has_groups  = input("فيه Group layer زي Elite/GC? (y/n): ").strip().lower() == "y"
        region_name = os.path.basename(region_path)
    else:
        print("❌ اختيار غلط"); input("\nEnter..."); return

    if not os.path.exists(region_path):
        print(f"\n❌ Path not found: {region_path}")
        input("\nEnter..."); return

    results = run(region_path, has_groups)

    out = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        f"campaigns_{region_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    export(results, out, region_name)
    print("\n✅ Done!")
    input("\nEnter للخروج...")

if __name__ == "__main__":
    main()